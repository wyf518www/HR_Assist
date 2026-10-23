"""Bounded document extraction. The HTTP process delegates parsing to a child."""
from pathlib import Path
import json
import re
import shutil
import subprocess
import sys
import tempfile

from .storage import BusinessError, validate_zip


def extract_file(path: Path):
    try:
        completed = subprocess.run(
            [sys.executable, "-m", "hr_assist.documents", str(path)],
            capture_output=True, timeout=180, encoding="utf-8", errors="replace",
        )
        result = json.loads(completed.stdout)
    except subprocess.TimeoutExpired as exc:
        raise BusinessError("PARSE_TIMEOUT", "文档解析超时，原件已保留。") from exc
    except (ValueError, OSError) as exc:
        raise BusinessError("PARSE_FAILED", "文档解析失败，原件已保留。") from exc
    if result.get("error"):
        raise BusinessError(result["code"], result["error"])
    return result["spans"]


def extract_inline(path):
    extension = path.suffix.lower()
    pages = []
    if extension == ".pdf":
        from pypdf import PdfReader
        if not path.read_bytes()[:1024].lstrip().startswith(b"%PDF-"):
            raise BusinessError("INVALID_PDF", "文件不是有效的 PDF。")
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise BusinessError("PDF_ENCRYPTED", "PDF 已加密，请提供未加密的原件。")
        if len(reader.pages) > 40:
            raise BusinessError("PAGE_LIMIT", "第一版每份文档最多处理 40 页。")
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            source = "text"
            if len(re.sub(r"\s", "", text)) < 25:
                if not shutil.which("tesseract"):
                    raise BusinessError("OCR_UNAVAILABLE", "此文档需要 OCR，但当前环境没有可用的 OCR 引擎。")
                import pypdfium2
                with tempfile.TemporaryDirectory() as temporary:
                    document = pypdfium2.PdfDocument(path)
                    try:
                        image = document[index].render(scale=2).to_pil()
                        image_path = Path(temporary) / "page.png"
                        image.save(image_path)
                    finally:
                        document.close()
                    task = subprocess.run(
                        ["tesseract", str(image_path), "stdout", "-l", "chi_sim+eng"],
                        capture_output=True, timeout=35, encoding="utf-8", errors="replace",
                    )
                    if task.returncode:
                        raise BusinessError("OCR_FAILED", "OCR 失败，请核对扫描页清晰度。")
                    text, source = task.stdout, "ocr"
            if len(text.strip()) < 10:
                raise BusinessError("EMPTY_PAGE", f"第 {index + 1} 页无法识别有效文字，请人工核验。")
            pages.append((index + 1, text, source))
    elif extension == ".docx":
        from docx import Document
        validate_zip(path.read_bytes())
        document = Document(path)
        texts = [p.text for p in document.paragraphs]
        for table in document.tables:
            texts.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        pages = [(None, "\n".join(texts), "docx")]
    elif extension == ".txt":
        pages = [(None, path.read_text(encoding="utf-8-sig"), "text")]
    else:
        raise BusinessError("UNSUPPORTED", "JD 支持 PDF、DOCX、UTF-8 TXT；简历只接受 PDF。")
    if sum(len(text) for _, text, _ in pages) > 100_000:
        raise BusinessError("TEXT_LIMIT", "文档文字过长，需要拆分后处理；没有截断后继续评分。")
    spans = []
    for page, text, source in pages:
        # Keep actual extraction text. No rephrasing of evidence.
        for paragraph in re.split(r"\n\s*\n", text):
            if paragraph.strip():
                spans.append({"id": f"R{len(spans)+1}", "page": page, "text": paragraph.strip(), "source": source})
    if not spans:
        raise BusinessError("EMPTY_DOCUMENT", "文档中没有可读取的文字。")
    return spans


if __name__ == "__main__":
    try:
        if sys.platform == "linux":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024**2, 1536 * 1024**2))
            resource.setrlimit(resource.RLIMIT_CPU, (150, 150))
            resource.setrlimit(resource.RLIMIT_FSIZE, (80 * 1024**2, 80 * 1024**2))
        value = {"spans": extract_inline(Path(sys.argv[1]))}
    except BusinessError as error:
        value = {"code": error.code, "error": error.message}
    except Exception:
        value = {"code": "PARSE_FAILED", "error": "文档损坏、编码不正确或解析器无法读取。"}
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(value, ensure_ascii=False))
