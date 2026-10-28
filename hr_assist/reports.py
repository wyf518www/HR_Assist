from copy import copy
from datetime import date, datetime
from io import BytesIO
import html
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .db import dump, now, uid
from .storage import BusinessError, validate_zip

LEGACY_MANUAL_COLUMNS=["推荐给部门日期","面试日期","简历是否通过用人部门筛选","面试评价","当前状态"]
# Keep the original five storage positions, including free-text notes and formulas.
MANUAL_COLUMNS=["推荐给用人部门日期","约面时间","用人部门复筛结果","面试评价","当前状态","面试状态","面试结果"]
MANUAL_CHOICES={"用人部门复筛结果":("通过","不通过"),"面试状态":("待面试","已面","拒绝面试"),"面试结果":("通过","待定","未通过")}
OVERVIEW_COLUMNS=["岗位名称","推荐日期","姓名","匹配度","推荐等级","当前公司","年限","核心优势","主要风险","建议",*MANUAL_COLUMNS]
FOLLOWUP_COLUMNS=["记录ID","岗位名称","姓名",*MANUAL_COLUMNS]


def empty_cells():
    return [{"value":None,"type":"n","format":"General"} for _ in MANUAL_COLUMNS]


def normalized_cells(cells):
    return cells + empty_cells()[len(cells):]


def validate_manual_cell(column, cell):
    value=cell_value(cell)
    if value is None or value == '':
        return
    if column in MANUAL_CHOICES and value not in MANUAL_CHOICES[column]:
        raise BusinessError('FOLLOWUP_VALUE', column+'只能填写：'+' / '.join(MANUAL_CHOICES[column])+'，也可留空。')
    if column in {'推荐给用人部门日期','约面时间'}:
        try:
            if cell['type']=='f' or not isinstance(value,(str,date,datetime)):
                raise ValueError()
            parsed=datetime.fromisoformat(value) if isinstance(value,str) else value
            if isinstance(parsed,datetime) and parsed.tzinfo is not None:
                raise ValueError()
        except ValueError as exc:
            raise BusinessError('FOLLOWUP_DATE', column+'请填写有效日期或北京时间（YYYY-MM-DD HH:MM），也可留空。') from exc


def serialize_cell(cell):
    value=cell.value
    if isinstance(value,datetime):
        value={"datetime":value.isoformat()}
    elif isinstance(value,date):
        value={"date":value.isoformat()}
    if isinstance(value,str) and len(value)>20_000:
        raise BusinessError("CELL_LIMIT", "人工反馈单元格过长。")
    return {"value":value,"type":cell.data_type,"format":cell.number_format}


def cell_value(cell):
    value=cell["value"]
    if isinstance(value,dict):
        if "datetime" in value:
            return datetime.fromisoformat(value["datetime"])
        return date.fromisoformat(value["date"])
    return value


def display_cell(cell):
    value=cell_value(cell)
    return value.isoformat() if isinstance(value,(date,datetime)) else value


def text_cell(cell, value):
    cell.value="" if value is None else str(value)
    cell.data_type="s"  # Do not interpret model/filename text as Excel formulas.


def write_manual(cell, original):
    cell.value=cell_value(original)
    cell.data_type=original["type"]
    cell.number_format=original["format"]


def style_sheet(sheet, manual_start):
    sheet.freeze_panes="D2"
    sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.font=Font(color="FFFFFF",bold=True)
        cell.fill=PatternFill("solid",fgColor="16766F" if cell.column>=manual_start else "23495C")
        cell.alignment=Alignment(wrap_text=True,vertical="center")
    sheet.row_dimensions[1].height=32
    for col in sheet.columns:
        letter=col[0].column_letter
        sheet.column_dimensions[letter].width=20 if col[0].column<8 else 38
        for cell in col[1:]:
            cell.alignment=Alignment(vertical="top",wrap_text=True)


def applications(db, *, identity=None, summary=False):
    # Lists need a few score fields, not every evidence quotation. Project in
    # SQLite so the full result never crosses into the response-building loop.
    result=("""CASE WHEN s.result IS NULL OR json_type(s.result)='null' THEN NULL
        WHEN json(s.result)='{}' THEN '{}'
        ELSE json_object('score',json_extract(s.result,'$.score'),
                         'grade',json_extract(s.result,'$.grade'),
                         'provisional',json_extract(s.result,'$.provisional')) END"""
        if summary else "s.result")
    where=" WHERE a.id=?" if identity is not None else ""
    rows=db.all(f"""SELECT a.*,r.name role_name,r.is_sales,{result} assessment_result,f.cells,f.revision
    FROM applications a LEFT JOIN roles r ON a.role_id=r.id
    LEFT JOIN assessments s ON a.assessment_id=s.id
    LEFT JOIN followups f ON f.application_id=a.id{where} ORDER BY a.created_at DESC,a.id""",
    (identity,) if identity is not None else ())
    for row in rows:
        row["assessment"]=json.loads(row.pop("assessment_result") or "null")
        if summary and row['assessment']:
            row['assessment']['provisional']=bool(row['assessment']['provisional'])
        row["manual_cells"]=normalized_cells(json.loads(row.pop("cells") or dump(empty_cells())))
        row["manual"]=dict(zip(MANUAL_COLUMNS,[display_cell(c) for c in row["manual_cells"]]))
    # Stable within date and role. ISO UTC date corresponds to stored local recommendation date.
    rows.sort(key=lambda x:(x["role_name"] or "待分类",x["id"]))
    rows.sort(key=lambda x:x["recommended_at"] or "",reverse=True)
    return rows


def application_counts(db):
    """Aggregate the overview without loading candidate evidence or feedback."""
    grades=("强烈推荐","推荐","待定","不推荐")
    grade_columns=','.join(f"coalesce(sum(json_extract(s.result,'$.grade')=?),0) grade_{i}" for i in range(len(grades)))
    row=db.one(f"""SELECT count(*) applications,
        coalesce(sum(CASE WHEN s.result IS NULL OR json_type(s.result)='null'
                         OR json(s.result)='{{}}' THEN 1 ELSE 0 END),0) pending,
        {grade_columns}
        FROM applications a LEFT JOIN assessments s ON a.assessment_id=s.id""",grades)
    return {'applications':row['applications'],'pending':row['pending'],
            'grades':{grade:row[f'grade_{i}'] for i,grade in enumerate(grades)}}


def save_followup(db, identity, data):
    with db.connect(True) as con:
        row=con.execute('SELECT cells,revision FROM followups WHERE application_id=?',(identity,)).fetchone()
        if not row:
            raise BusinessError('APPLICATION_NOT_FOUND','申请记录不存在。',404)
        if row['revision']!=data.expected_revision:
            raise BusinessError('FEEDBACK_CONFLICT','这条人工跟进已被其他页面或 Excel 修改，本次未保存。请刷新并核对最新内容。',409)
        cells=normalized_cells(json.loads(row['cells']))
        changed=False
        for column,value in data.changes.items():
            index=MANUAL_COLUMNS.index(column)
            cell={'value':value or None,'type':'s' if value else 'n','format':'General'}
            validate_manual_cell(column,cell)
            if display_cell(cells[index]) != (value or None):
                cells[index]=cell
                changed=True
        if changed:
            con.execute('UPDATE followups SET cells=?,revision=revision+1,updated_at=? WHERE application_id=?',
                        (dump(cells),now(),identity))
    if changed:
        db.event('FOLLOWUP_SAVED','人工跟进字段已保存。',application_id=identity)
    return {'ok':True,'revision':row['revision']+int(changed)}


def save_followup_batch(db, data):
    # A patch is explicit: omitted columns are never inspected or rewritten,
    # including historical free-text values, formulas and cell formatting.
    patch={}
    for column,value in data.fields.items():
        cell={'value':value or None,'type':'s' if value else 'n','format':'General'}
        validate_manual_cell(column,cell)
        patch[MANUAL_COLUMNS.index(column)]=cell
    changed_rows={}
    revisions={}
    with db.connect(True) as con:
        for item in data.items:
            row=con.execute('SELECT cells,revision FROM followups WHERE application_id=?',(item.application_id,)).fetchone()
            if row is None:
                raise BusinessError('APPLICATION_NOT_FOUND','部分申请记录不存在，本次全部未保存，请刷新候选人列表。',404)
            if row['revision']!=item.expected_revision:
                raise BusinessError('FEEDBACK_CONFLICT','所选记录的人工跟进已被其他页面或 Excel 修改，本次全部未保存。请刷新并核对最新内容。',409)
            cells=normalized_cells(json.loads(row['cells']))
            changed=False
            for index,cell in patch.items():
                if display_cell(cells[index])!=cell['value']:
                    cells[index]=cell
                    changed=True
            revisions[item.application_id]=row['revision']+int(changed)
            if changed:
                changed_rows[item.application_id]=cells
        # Validate every selected revision before the first write.
        timestamp=now()
        for identity,cells in changed_rows.items():
            con.execute('UPDATE followups SET cells=?,revision=revision+1,updated_at=? WHERE application_id=?',
                        (dump(cells),timestamp,identity))
    if changed_rows:
        db.event('FOLLOWUP_BATCH_SAVED',f'人工跟进批量保存：选中 {len(data.items)} 条，更新 {len(changed_rows)} 条。')
    return {'ok':True,'selected_rows':len(data.items),'updated_rows':len(changed_rows),'revisions':revisions}


def followup_export(db, storage):
    rows=applications(db)
    export_id=uid()
    book=Workbook()
    sheet=book.active
    sheet.title="人工跟进"
    sheet.append(FOLLOWUP_COLUMNS)
    snapshot={}
    for row_index,row in enumerate(rows,2):
        for column,value in enumerate((row["id"],row["role_name"] or "待分类",row["name"]),1):
            text_cell(sheet.cell(row_index,column),value)
        for column,cell in enumerate(row["manual_cells"],4):
            write_manual(sheet.cell(row_index,column),cell)
        snapshot[row["id"]]=row["manual_cells"]
    metadata=book.create_sheet("_metadata")
    metadata.append(["template","HR_ASSIST_FOLLOWUP_V2"])
    metadata.append(["export_id",export_id])
    metadata.sheet_state="veryHidden"
    style_sheet(sheet,4)
    sheet.column_dimensions["A"].width=36
    for name,choices in MANUAL_CHOICES.items():
        column=get_column_letter(4+MANUAL_COLUMNS.index(name))
        validation=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=True)
        validation.showErrorMessage=True
        validation.errorTitle='请选择有效选项'
        validation.error='请从下拉列表选择，也可留空。'
        sheet.add_data_validation(validation)
        validation.add(f'{column}2:{column}{max(2,len(rows)+1)}')
    buffer=BytesIO();book.save(buffer)
    file=storage.put(buffer.getvalue(),"人工跟进.xlsx","followup")
    db.execute("INSERT INTO workbook_exports VALUES (?,?,?,?)",(export_id,dump(snapshot),file["id"],now()))
    return file


def followup_import(db, storage, data, filename):
    validate_zip(data)
    try:
        book=load_workbook(BytesIO(data),data_only=False,keep_links=False)
        meta=book["_metadata"]
        template=meta["B1"].value
        if template not in {"HR_ASSIST_FOLLOWUP_V1","HR_ASSIST_FOLLOWUP_V2"}:
            raise ValueError()
        columns=FOLLOWUP_COLUMNS if template=='HR_ASSIST_FOLLOWUP_V2' else ['记录ID','岗位名称','姓名',*LEGACY_MANUAL_COLUMNS]
        export_id=meta["B2"].value
        source=db.one("SELECT * FROM workbook_exports WHERE id=?",(export_id,))
        if not source:
            raise BusinessError("UNKNOWN_EXPORT","此文件不是当前工作台导出的人工跟进表。")
        sheet=book["人工跟进"]
        if sheet.max_row>20_001 or sheet.max_column!=len(columns):
            raise ValueError()
        if [cell.value for cell in sheet[1]]!=columns:
            raise ValueError()
        baseline=json.loads(source["snapshot"])
        incoming={}
        for row in sheet.iter_rows(min_row=2):
            identity=row[0].value
            if identity is None and all(c.value is None for c in row):
                continue
            if identity not in baseline or identity in incoming:
                raise BusinessError("ROW_ID_INVALID","记录 ID 被改写、重复或来自其他导出。请重新下载模板。")
            incoming[identity]=[serialize_cell(c) for c in row[3:len(columns)]]
    except BusinessError:
        raise
    except Exception as exc:
        raise BusinessError("WORKBOOK_INVALID","请上传工作台导出的完整 .xlsx；不要修改表头和隐藏元数据。") from exc
    changes={}
    conflicts=[]
    with db.connect(True) as con:
        for identity,cells in incoming.items():
            current_row=con.execute("SELECT cells,revision FROM followups WHERE application_id=?",(identity,)).fetchone()
            if not current_row:
                raise BusinessError("ROW_REMOVED","申请记录已不存在，无法合并。",409)
            current=normalized_cells(json.loads(current_row["cells"]))
            merged=copy(current)
            for i,cell in enumerate(cells):
                before=normalized_cells(baseline[identity])[i]
                if cell!=before:
                    if template=='HR_ASSIST_FOLLOWUP_V2':
                        validate_manual_cell(MANUAL_COLUMNS[i],cell)
                    if current[i]!=before and current[i]!=cell:
                        conflicts.append({"application_id":identity,"column":MANUAL_COLUMNS[i]})
                    else:
                        merged[i]=cell
            if merged!=current:
                changes[identity]=(merged,current_row["revision"]+1)
        if conflicts:
            raise BusinessError("FEEDBACK_CONFLICT",f"发现 {len(conflicts)} 个字段与后来保存的反馈冲突。本次未写入，请重新下载最新表后合并。",409)
        for identity,(cells,revision) in changes.items():
            con.execute("UPDATE followups SET cells=?,revision=?,updated_at=? WHERE application_id=?",(dump(cells),revision,now(),identity))
        # Original upload is immutable. Transactions cover feedback updates and the import record.
        # File registration happens after conflict validation; its orphan is harmless if DB commit fails.
        file_id=uid()
        import hashlib
        from .storage import safe_filename
        relative=f"followup/{file_id}_{safe_filename(filename)}"
        path=storage.path(relative)
        with path.open("xb") as stream:
            stream.write(data)
        con.execute("INSERT INTO files VALUES (?,?,?,?,?,?,?)",(file_id,hashlib.sha256(data).hexdigest(),"followup",filename,relative,len(data),now()))
        result={"updated_rows":len(changes),"received_rows":len(incoming),"omitted_rows_preserved":len(baseline)-len(incoming)}
        con.execute("INSERT INTO feedback_imports VALUES (?,?,?,?,?)",(uid(),export_id,file_id,dump(result),now()))
    db.event("FOLLOWUP_IMPORTED",f"人工跟进导入完成：更新 {len(changes)} 条。")
    return result


def overview_export(db, storage):
    rows=applications(db)
    book=Workbook()
    sheet=book.active
    sheet.title="招聘推荐总览"
    sheet.append([*OVERVIEW_COLUMNS,"_application_id"])
    for index,row in enumerate(rows,2):
        result=row["assessment"] or {}
        refs=[]
        for item in result.get("items",[]):
            citations="、".join(e["span_id"] for e in item["evidence"])
            if citations:
                refs.append(f"{item['description']} [{citations}]：{item['reason']}")
        facts=result.get("facts",{})
        values=[row["role_name"] or "待分类",row["recommended_at"] or "",row["name"],result.get("score"),result.get("grade",""),
                facts.get("company",""),facts.get("experience",""),"；".join(result.get("advantages",[])),
                "；".join(result.get("risks",[])) or ("JD缺失项：岗位要求未补齐或尚未确认，停止评估" if row["status"]=="WAITING_JD" else row["status"] if not result else result.get("gate_reason","")),
                ("暂计分；存在未知信息，请人工核验。\n" if result.get("provisional") else "")+("\n".join(refs) or ("请查看逐项要求与待核验信息。" if result else "尚未评估，请在工作台处理待办。"))]
        for column,value in enumerate(values,1):
            if column==4:
                sheet.cell(index,column,value)
            else:
                text_cell(sheet.cell(index,column),value)
        for column,cell in enumerate(row["manual_cells"],11):
            write_manual(sheet.cell(index,column),cell)
        text_cell(sheet.cell(index,len(OVERVIEW_COLUMNS)+1),row["id"])
    sheet.column_dimensions[get_column_letter(len(OVERVIEW_COLUMNS)+1)].hidden=True
    style_sheet(sheet,11)
    evidence=book.create_sheet("评估证据")
    evidence.append(["记录ID","岗位","JD版本","维度","要求","判定","得分","简历位置","简历引文","JD依据来源"])
    for row in rows:
        result=row["assessment"] or {}
        span_map={s["id"]:s for s in result.get("evidence_spans",[])}
        for item in result.get("items",[]):
            empty_quote="已确认JD明确不作要求，无须简历证据" if item["verdict"]=="met" and item.get("hard") else "无证据，待核验"
            for ref in item["evidence"] or [{"span_id":"","quote":empty_quote}]:
                span=span_map.get(ref["span_id"],{})
                values=[row["id"],row["role_name"],result.get("jd_version_id",""),item["dimension"],item["description"],item["verdict"],item["awarded"],
                        f"{ref['span_id']} / 第 {span.get('page') or '—'} 页",ref["quote"],"已确认 JD 结构化要求" if item["source"]=="jd" else "用户确认的全局评分规则"]
                for col,value in enumerate(values,1):
                    text_cell(evidence.cell(evidence.max_row+1 if col==1 else evidence.max_row,col),value)
    style_sheet(evidence,99)
    buffer=BytesIO();book.save(buffer)
    file=storage.put(buffer.getvalue(),"招聘推荐总览.xlsx","reports")
    db.put("latest_overview",file["id"])
    return file


def role_reports(db, storage, run_id):
    groups={}
    for row in applications(db):
        if row["assessment"]:
            groups.setdefault(row["role_name"],[]).append(row)
    files=[]
    escape=html.escape
    for role,rows in groups.items():
        sections=[]
        for row in rows:
            result=row["assessment"]
            items=[]
            for item in result["items"]:
                quotes="<br>".join(escape(f"{e['span_id']}：{e['quote']}") for e in item["evidence"])
                verdict={"met":"满足","partial":"部分满足","not_met":"不满足","unknown":"未知，待核验"}[item["verdict"]]
                items.append(f"<tr><td>{escape(item['description'])}</td><td>{verdict}</td><td>{item['awarded']}/{item['points']}</td><td>{escape(item['reason'])}<br>{quotes}</td></tr>")
            provisional="<p><b>暂计分；存在未知信息，请人工核验。缺少证据不等于不满足要求。</b></p>" if result.get("provisional") else ""
            sections.append(f"<section><h2>{escape(row['name'])} · {result['score']}/100 · {escape(result['grade'])}</h2>{provisional}<p>{escape(result['gate_reason'])}</p><p>记录 {row['id']} / JD {escape(result.get('jd_version_id',''))}</p><table><tr><th>要求</th><th>判定</th><th>分数</th><th>依据</th></tr>{''.join(items)}</table></section>")
        document=f"<!doctype html><html lang='zh-CN'><meta charset='utf-8'><meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; style-src 'unsafe-inline'\"><title>{escape(role)}推荐报告</title><style>body{{max-width:1100px;margin:40px auto;padding:20px;font:15px/1.8 sans-serif;color:#183443}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ccd7dc;padding:12px;text-align:left}}section{{margin-top:35px}}</style><h1>{escape(role)}推荐报告</h1><p>生成时间：{now()} / 批次：{run_id}。建议供人工核验，不自动淘汰或录用。</p>{''.join(sections)}</html>"
        files.append(storage.put(document.encode(),f"{role}_推荐报告.html","reports"))
    return files
