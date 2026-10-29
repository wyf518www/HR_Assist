"""Spreadsheet display escapes must not alter recruitment source records."""
from copy import deepcopy
from io import BytesIO
import re

import pytest
from openpyxl import Workbook, load_workbook

from hr_assist.db import dump, now, uid
from hr_assist.reports import empty_cells, excel_display_text, followup_export, overview_export, text_cell


UNSUPPORTED = [code for code in range(32) if code not in {9, 10, 13}] + [0xD800, 0xDFFF, 0xFFFE, 0xFFFF]


@pytest.mark.parametrize('code', UNSUPPORTED)
def test_xml_unsupported_codepoints_are_visible_and_excel_roundtrip_is_safe(code):
    book = Workbook()
    source = '=合成资料' + chr(code) + '后缀'
    cell = book.active['A1']
    text_cell(cell, source)
    output = BytesIO()
    book.save(output)
    actual = load_workbook(BytesIO(output.getvalue())).active['A1']
    assert actual.value == '=合成资料' + f'[U+{code:04X}]' + '后缀'
    assert actual.data_type == 's'
    assert '原始简历' in actual.comment.text and '未修改' in actual.comment.text
    assert '合成资料' not in actual.comment.text
    assert source == '=合成资料' + chr(code) + '后缀'


def test_normal_unicode_and_tab_cr_lf_roundtrip_without_display_change():
    source = '=中文\tEnglish\n换行\r回车\r\n𠀀🙂\u200dé'
    book = Workbook()
    text_cell(book.active['A1'], source)
    output = BytesIO()
    book.save(output)
    actual = load_workbook(BytesIO(output.getvalue())).active['A1']
    assert actual.value == source and actual.data_type == 's'
    assert actual.comment is None


@pytest.mark.parametrize('source', ['=自动汇总' + '\x0b' * 6000, '🙂' * 20000],
                         ids=['escaped-control-expansion', 'supplementary-unicode'])
def test_long_automatic_summary_is_explicitly_shortened_without_silent_excel_truncation(source):
    book = Workbook()
    text_cell(book.active['A1'], source)
    output = BytesIO()
    book.save(output)
    actual = load_workbook(BytesIO(output.getvalue())).active['A1']
    suffix = '\n[内容过长，已缩短显示；完整内容请在工作台查看]'
    assert actual.value.endswith(suffix)
    assert len(actual.value.encode('utf-16-le')) // 2 <= 32767
    prefix = actual.value[:-len(suffix)]
    assert excel_display_text(source).startswith(prefix)
    assert '[U+' not in re.sub(r'\[U\+[0-9A-F]{4,6}\]', '', prefix)
    assert actual.data_type == 's'
    assert '仅缩短此文件中的显示' in actual.comment.text
    assert '完整内容保留在工作台' in actual.comment.text


def seed(service, application, role, source):
    ref = {'span_id': 'R1', 'quote': source, 'semantic': True, 'interpretation': '模型\x0c解读', 'page': 1, 'source': 'text'}
    result = {'score': 90, 'grade': '推荐', 'provisional': False, 'gate_reason': '合成政策说明',
              'jd_version_id': role['current_version'], 'facts': {'company': '=示例\x00公司'},
              'advantages': ['优势\x0b说明'], 'risks': [],
              'items': [{'id': 'skills', 'dimension': 'skills', 'description': '要求\x1f内容',
                         'verdict': 'met', 'awarded': 30, 'points': 30, 'reason': '理由\x01内容',
                         'source': 'jd', 'hard': False, 'evidence': [ref]}],
              'evidence_spans': [{'id': 'R1', 'text': source, 'page': 1, 'source': 'text'}]}
    identity = uid()
    service.db.execute('INSERT INTO assessments VALUES (?,?,?,?,?,?,?)',
                       (identity, application['id'], role['current_version'], uid(), 'COMPLETED', dump(result), now()))
    service.db.execute("UPDATE applications SET assessment_id=?,status='ASSESSED',name=? WHERE id=?",
                       (identity, '合成\x0b候选人', application['id']))
    cells = empty_cells()
    cells[3] = {'value': '=1+1', 'type': 'f', 'format': '0.00'}
    cells[4] = {'value': '人工\t备注\n第二行', 'type': 's', 'format': '@'}
    service.db.execute('UPDATE followups SET cells=?,revision=3 WHERE application_id=?', (dump(cells), application['id']))
    return identity


def test_overview_exports_control_characters_without_modifying_source_or_human_cells(service, application, role):
    source = '=原文\x0c控制\x00字符\t和正常\n换行'
    identity = seed(service, application, role, source)
    before = {
        'assessment': service.db.one('SELECT * FROM assessments WHERE id=?', (identity,)),
        'followup': service.db.one('SELECT * FROM followups WHERE application_id=?', (application['id'],)),
        'application': service.db.one('SELECT * FROM applications WHERE id=?', (application['id'],)),
        'documents': service.db.all('SELECT * FROM documents'),
        'pdf': service.storage.file(application['file_id'])[1].read_bytes(),
    }
    snapshot = deepcopy(before)
    exported = overview_export(service.db, service.storage)
    book = load_workbook(service.storage.file(exported['id'])[1], data_only=False)
    overview, evidence = book['招聘推荐总览'], book['评估证据']
    assert overview['C2'].value == '合成[U+000B]候选人'
    assert overview['F2'].value == '=示例[U+0000]公司' and overview['F2'].data_type == 's'
    assert evidence['I2'].value == excel_display_text(source) and evidence['I2'].data_type == 's'
    assert evidence['K2'].value == '模型[U+000C]解读'
    assert evidence['I2'].comment and evidence['K2'].comment
    assert overview['N2'].value == '=1+1' and overview['N2'].data_type == 'f' and overview['N2'].number_format == '0.00'
    assert overview['O2'].value == '人工\t备注\n第二行' and overview['O2'].number_format == '@'
    assert overview['N2'].comment is None and overview['O2'].comment is None
    assert service.db.one('SELECT * FROM assessments WHERE id=?', (identity,)) == snapshot['assessment']
    assert service.db.one('SELECT * FROM followups WHERE application_id=?', (application['id'],)) == snapshot['followup']
    assert service.db.one('SELECT * FROM applications WHERE id=?', (application['id'],)) == snapshot['application']
    assert service.db.all('SELECT * FROM documents') == snapshot['documents']
    assert service.storage.file(application['file_id'])[1].read_bytes() == snapshot['pdf']
    assert before == snapshot


def test_evidence_is_split_after_escaping_without_lost_text_or_broken_markers(service, application, role):
    source = '甲' * 14997 + '\x0c' + '乙' * 3 + '\x00' * 16000 + '尾部'
    seed(service, application, role, source)
    exported = overview_export(service.db, service.storage)
    sheet = load_workbook(service.storage.file(exported['id'])[1])['评估证据']
    chunks = [sheet.cell(index, 9).value for index in range(2, sheet.max_row + 1)]
    assert len(chunks) > 3
    assert ''.join(chunks) == excel_display_text(source)
    assert all(0 < len(chunk) <= 15000 for chunk in chunks)
    assert all('[U+' not in re.sub(r'\[U\+[0-9A-F]{4,6}\]', '', chunk) for chunk in chunks)
    assert all(sheet.cell(index, 9).comment for index in range(2, sheet.max_row + 1))
    assert all(f'{index-1}/{len(chunks)}' in sheet.cell(index, 8).value for index in range(2, sheet.max_row + 1))


def test_followup_export_escapes_automatic_identity_only_and_keeps_human_formula(service, application, role):
    seed(service, application, role, '合成原文')
    before = service.db.one('SELECT * FROM followups WHERE application_id=?', (application['id'],))
    exported = followup_export(service.db, service.storage)
    sheet = load_workbook(service.storage.file(exported['id'])[1], data_only=False)['人工跟进']
    assert sheet['C2'].value == '合成[U+000B]候选人' and sheet['C2'].comment
    assert sheet['G2'].value == '=1+1' and sheet['G2'].data_type == 'f' and sheet['G2'].number_format == '0.00'
    assert sheet['H2'].value == '人工\t备注\n第二行' and sheet['H2'].number_format == '@'
    assert service.db.one('SELECT * FROM followups WHERE application_id=?', (application['id'],)) == before
