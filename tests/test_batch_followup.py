from io import BytesIO
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
import pytest

from hr_assist.db import dump
from hr_assist.main import create_app
from hr_assist.reports import (MANUAL_COLUMNS, applications, empty_cells,
                               save_followup, save_followup_batch)
from hr_assist.schemas import FollowupBatchInput, FollowupInput
from hr_assist.storage import BusinessError
from conftest import make_pdf
from test_excel import exported, upload


def request_for(ids,fields,revision=0):
    return FollowupBatchInput(confirmed=True,
        items=[{'application_id':identity,'expected_revision':revision} for identity in ids],fields=fields)


def all_followups(service):
    return service.db.all('SELECT * FROM followups ORDER BY application_id')


def test_batch_patch_preserves_unselected_records_legacy_cells_and_identity(service,application):
    second=service.ingest(make_pdf('Second synthetic resume.'),'second.pdf')
    unselected=service.ingest(make_pdf('Unselected synthetic resume.'),'unselected.pdf')
    legacy=empty_cells()
    legacy[2]={'value':'旧版待部门意见','type':'s','format':'@'}
    legacy[3]={'value':'=1+1','type':'f','format':'0.00'}
    legacy[4]={'value':'独立的人工备注','type':'s','format':'@'}
    service.db.execute('UPDATE followups SET cells=? WHERE application_id=?',(dump(legacy),application['id']))
    previous_apps=service.db.all('SELECT * FROM applications ORDER BY id')
    previous_unselected=service.db.one('SELECT * FROM followups WHERE application_id=?',(unselected['id'],))
    result=save_followup_batch(service.db,request_for([application['id'],second['id']],
        {'约面时间':'2026-10-29T14:30','面试状态':'待面试'}))
    assert result=={'ok':True,'selected_rows':2,'updated_rows':2,
                    'revisions':{application['id']:1,second['id']:1}}
    rows={row['id']:row for row in applications(service.db)}
    for identity in (application['id'],second['id']):
        assert rows[identity]['manual']['约面时间']=='2026-10-29T14:30'
        assert rows[identity]['manual']['面试状态']=='待面试'
    assert rows[application['id']]['manual_cells'][2:5]==legacy[2:5]
    assert service.db.all('SELECT * FROM applications ORDER BY id')==previous_apps
    assert service.db.one('SELECT * FROM followups WHERE application_id=?',(unselected['id'],))==previous_unselected
    event=service.db.one("SELECT message FROM events WHERE code='FOLLOWUP_BATCH_SAVED'")['message']
    assert '2' in event and '14:30' not in event and 'Test' not in event


def test_explicit_clear_and_noop_keep_other_values_and_revisions(service,application):
    identity=application['id']
    save_followup(service.db,identity,FollowupInput(expected_revision=0,
        changes={'约面时间':'2026-10-29T14:30','面试状态':'待面试','面试结果':'待定'}))
    result=save_followup_batch(service.db,request_for([identity],{'约面时间':'','面试结果':None},1))
    row=applications(service.db)[0]
    assert result['updated_rows']==1 and row['revision']==2
    assert row['manual']['约面时间'] is None and row['manual']['面试结果'] is None
    assert row['manual']['面试状态']=='待面试'
    before=all_followups(service)
    result=save_followup_batch(service.db,request_for([identity],{'约面时间':'','面试状态':'待面试'},2))
    assert result['updated_rows']==0 and result['revisions'][identity]==2
    assert all_followups(service)==before


def test_batch_free_text_is_plain_text_in_excel_and_bounded(service,application):
    save_followup_batch(service.db,request_for([application['id']],
        {'面试评价':'=SUM(1,2)','当前状态':'统一备注\n等待本人回复'}))
    book=load_workbook(BytesIO(exported(service)))
    sheet=book['人工跟进']
    assert sheet['G2'].value=='=SUM(1,2)' and sheet['G2'].data_type=='s'
    assert sheet['H2'].value=='统一备注\n等待本人回复'
    with pytest.raises(ValidationError):
        request_for([application['id']],{'面试评价':'x'*20_001})


@pytest.mark.parametrize('missing',[False,True])
def test_batch_conflict_or_missing_row_rolls_back_all(service,application,missing):
    second=service.ingest(make_pdf('Second synthetic resume.'),'second.pdf')
    save_followup(service.db,second['id'],FollowupInput(expected_revision=0,changes={'面试状态':'已面'}))
    before=all_followups(service)
    ids=[application['id'],'missing-record' if missing else second['id']]
    with pytest.raises(BusinessError) as error:
        save_followup_batch(service.db,request_for(ids,{'面试状态':'待面试'}))
    assert error.value.code==('APPLICATION_NOT_FOUND' if missing else 'FEEDBACK_CONFLICT')
    assert all_followups(service)==before
    assert not service.db.one("SELECT 1 FROM events WHERE code='FOLLOWUP_BATCH_SAVED'")


@pytest.mark.parametrize('column,value',[
    ('用人部门复筛结果','待定'),('面试状态','录用'),('面试结果','优秀'),
    ('约面时间','2026-02-30T12:00'),('推荐给用人部门日期','yesterday'),
    ('约面时间','2026-10-29T01:00+00:00'),('约面时间','=1+1')])
def test_batch_invalid_values_never_write(service,application,column,value):
    before=all_followups(service)
    with pytest.raises(BusinessError):
        save_followup_batch(service.db,request_for([application['id']],{column:value}))
    assert all_followups(service)==before


@pytest.mark.parametrize('changes',[
    {'confirmed':False},{'items':[]},{'fields':{}},
    {'items':[{'application_id':'a','expected_revision':0}]*2},
    {'items':[{'application_id':str(i),'expected_revision':0} for i in range(501)]},
    {'fields':{'score':'100'}},{'fields':{'面试状态':'x'*41}},
    {'items':[{'application_id':'a','expected_revision':-1}]},
])
def test_batch_input_contract(changes):
    body={'confirmed':True,'items':[{'application_id':'a','expected_revision':0}],'fields':{'面试状态':'待面试'}}
    with pytest.raises(ValidationError):
        FollowupBatchInput.model_validate({**body,**changes})


def test_excel_merge_after_batch_preserves_patch_and_rejects_conflict(service,application):
    baseline=exported(service)
    identity=application['id']
    save_followup_batch(service.db,request_for([identity],{'面试状态':'待面试','约面时间':'2026-10-29T14:30'}))
    # An old workbook editing another field can merge while keeping batch data.
    book=load_workbook(BytesIO(baseline));book['人工跟进']['F2']='通过'
    buffer=BytesIO();book.save(buffer)
    assert upload(service,buffer.getvalue())['updated_rows']==1
    row=applications(service.db)[0]
    assert row['revision']==2 and row['manual']['面试状态']=='待面试'
    assert row['manual']['用人部门复筛结果']=='通过'
    # The same old workbook cannot overwrite a field changed by this batch.
    before=all_followups(service)
    book['人工跟进']['I2']='已面';buffer=BytesIO();book.save(buffer)
    with pytest.raises(BusinessError) as error:
        upload(service,buffer.getvalue())
    assert error.value.code=='FEEDBACK_CONFLICT' and all_followups(service)==before
    with pytest.raises(BusinessError) as error:
        save_followup_batch(service.db,request_for([identity],{'面试结果':'通过'},1))
    assert error.value.code=='FEEDBACK_CONFLICT' and all_followups(service)==before


def test_batch_api_requires_auth_csrf_confirmation_and_returns_revisions(environment):
    with TestClient(create_app(environment)) as client:
        service=client.app.state.service
        first=service.ingest(make_pdf(),'first.pdf')
        second=service.ingest(make_pdf('Second synthetic resume.'),'second.pdf')
        body=request_for([first['id'],second['id']],{'面试状态':'待面试'}).model_dump()
        url='/api/applications/followup-batch'
        assert client.put(url,json=body).status_code==401
        session=client.post('/api/auth/login',json={'username':'管理员','password':environment.admin_password}).json()
        assert client.put(url,json=body).status_code==403
        client.headers['X-CSRF-Token']=session['csrf']
        assert client.put(url,json={k:v for k,v in body.items() if k!='confirmed'}).status_code==422
        assert client.put(url,json={**body,'confirmed':False}).status_code==422
        assert client.put(url,json={**body,'fields':{'score':'100'}}).status_code==422
        result=client.put(url,json=body)
        assert result.status_code==200 and result.json()['updated_rows']==2
        assert result.json()['revisions']=={first['id']:1,second['id']:1}
        assert client.put(url,json=body).status_code==409
        latest=client.get('/api/applications/'+first['id']).json()
        assert latest['manual']['面试状态']=='待面试' and latest['revision']==1
