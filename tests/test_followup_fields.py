from datetime import datetime
from io import BytesIO
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from hr_assist.db import dump
from hr_assist.main import create_app
from hr_assist.reports import (MANUAL_COLUMNS, LEGACY_MANUAL_COLUMNS, applications, empty_cells,
                               followup_import, overview_export, save_followup)
from hr_assist.schemas import FollowupInput
from hr_assist.storage import BusinessError
from conftest import make_pdf
from test_excel import exported, edit, upload


def save(service,identity,changes,revision=0):
    return save_followup(service.db,identity,FollowupInput(changes=changes,expected_revision=revision))


def test_all_new_fields_roundtrip_ui_excel_report(service,application):
    values={'推荐给用人部门日期':'2026-10-24','约面时间':'2026-10-26T09:30',
            '用人部门复筛结果':'通过','面试状态':'待面试','面试结果':'待定'}
    save(service,application['id'],values)
    row=applications(service.db)[0]
    assert row['revision']==1 and all(row['manual'][k]==v for k,v in values.items())
    workbook=exported(service)
    book=load_workbook(BytesIO(workbook));sheet=book['人工跟进']
    assert book['_metadata']['B1'].value=='HR_ASSIST_FOLLOWUP_V2'
    assert [c.value for c in sheet[1]]==['记录ID','岗位名称','姓名',*MANUAL_COLUMNS]
    assert len(sheet.data_validations.dataValidation)==3
    assert upload(service,workbook)['updated_rows']==0
    upload(service,edit(workbook,{'I2':'已面','J2':'通过','E2':datetime(2026,10,26,14,45)}))
    row=applications(service.db)[0]
    assert row['manual']['约面时间']=='2026-10-26T14:45:00'
    assert row['manual']['面试状态']=='已面' and row['manual']['面试结果']=='通过'
    report=overview_export(service.db,service.storage)
    sheet=load_workbook(service.storage.file(report['id'])[1])['招聘推荐总览']
    assert sheet['P2'].value=='已面' and sheet['Q2'].value=='通过'
    assert sheet['L2'].value==datetime(2026,10,26,14,45)


def test_legacy_v1_import_preserves_new_fields_and_original_notes(service,application):
    old_cells=empty_cells()[:5]
    old_cells[3]={'value':'=1+1','type':'f','format':'General'}
    old_cells[4]={'value':'旧自由备注','type':'s','format':'@'}
    service.db.execute('UPDATE followups SET cells=?',(dump(old_cells),))
    workbook=exported(service)
    book=load_workbook(BytesIO(workbook));sheet=book['人工跟进']
    sheet.delete_cols(9,2)
    for i,value in enumerate(['记录ID','岗位名称','姓名',*LEGACY_MANUAL_COLUMNS],1):
        sheet.cell(1,i,value)
    book['_metadata']['B1']='HR_ASSIST_FOLLOWUP_V1'
    export_id=book['_metadata']['B2'].value
    service.db.execute('UPDATE workbook_exports SET snapshot=? WHERE id=?',(dump({application['id']:old_cells}),export_id))
    sheet['F2']='已通过'
    output=BytesIO();book.save(output)
    save(service,application['id'],{'面试状态':'待面试','面试结果':'待定'})
    assert upload(service,output.getvalue())['updated_rows']==1
    cells=json.loads(service.db.one('SELECT cells FROM followups')['cells'])
    assert cells[3:5]==old_cells[3:5]
    row=applications(service.db)[0]
    assert row['manual']['用人部门复筛结果']=='已通过'
    assert row['manual']['面试状态']=='待面试' and row['manual']['面试结果']=='待定'
    # Legacy freeform enum values are retained in V2 until explicitly edited.
    assert upload(service,exported(service))['updated_rows']==0


def test_web_excel_conflict_is_atomic_and_nonoverlapping_merges(service,application):
    workbook=exported(service)
    save(service,application['id'],{'面试状态':'已面'})
    upload(service,edit(workbook,{'F2':'通过'}))
    assert applications(service.db)[0]['revision']==2
    before=service.db.one('SELECT * FROM followups')
    with pytest.raises(BusinessError) as error:
        upload(service,edit(workbook,{'I2':'拒绝面试','J2':'通过'}))
    assert error.value.code=='FEEDBACK_CONFLICT'
    assert service.db.one('SELECT * FROM followups')==before
    with pytest.raises(BusinessError) as error:
        save(service,application['id'],{'面试结果':'未通过'},revision=0)
    assert error.value.status==409 and service.db.one('SELECT * FROM followups')==before


@pytest.mark.parametrize('column,value',[
    ('用人部门复筛结果','待定'),('面试状态','录用'),('面试结果','优秀'),
    ('约面时间','2026-02-30T12:00'),('推荐给用人部门日期','yesterday'),
    ('约面时间','2026-10-25T01:00+00:00'),('约面时间','=1+1')])
def test_invalid_manual_updates_do_not_write(service,application,column,value):
    before=service.db.one('SELECT * FROM followups')
    with pytest.raises(BusinessError):
        save(service,application['id'],{column:value})
    assert service.db.one('SELECT * FROM followups')==before
    position=chr(ord('D')+MANUAL_COLUMNS.index(column))+'2'
    with pytest.raises(BusinessError):
        upload(service,edit(exported(service),{position:value}))
    assert service.db.one('SELECT * FROM followups')==before


def test_followup_api_auth_csrf_and_conflict(environment):
    with TestClient(create_app(environment)) as client:
        service=client.app.state.service
        app=service.ingest(make_pdf(),'【TestRole_Test】.pdf')
        url='/api/applications/'+app['id']+'/followup'
        body={'expected_revision':0,'changes':{'面试状态':'待面试'}}
        assert client.put(url,json=body).status_code==401
        session=client.post('/api/auth/login',json={'username':'管理员','password':environment.admin_password}).json()
        assert client.put(url,json=body).status_code==403
        client.headers['X-CSRF-Token']=session['csrf']
        assert client.put(url,json={**body,'changes':{'score':'100'}}).status_code==422
        assert client.put(url,json=body).status_code==200
        assert client.put(url,json=body).status_code==409
        latest=client.get('/api/applications/'+app['id']).json()
        assert latest['manual']['面试状态']=='待面试' and latest['revision']==1
