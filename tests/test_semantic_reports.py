from openpyxl import load_workbook

from hr_assist.db import uid, dump, now
from hr_assist.reports import overview_export, role_reports


def seed_score(service,application,role,semantic=True):
    source='=原始资料 <原文>\n'+'原始正文'*9000
    ref={'span_id':'R1','quote':source}
    if semantic:ref.update(semantic=True,interpretation='模型对相关能力的概括 <解读>',page=1,source='text')
    result={'score':90,'grade':'推荐','provisional':False,'gate_reason':'合成政策说明',
        'jd_version_id':role['current_version'],'facts':{},'advantages':[],'risks':[],
        'items':[{'id':'skills','dimension':'skills','description':'合成要求','verdict':'met','awarded':30,
                  'points':30,'reason':'合成理由','source':'jd','hard':False,'evidence':[ref]}],
        'evidence_spans':[{'id':'R1','page':1,'text':source,'source':'text'}]}
    if semantic:result['evidence_mode']='model_semantic_review'
    sid=uid()
    service.db.execute('INSERT INTO assessments VALUES (?,?,?,?,?,?,?)',(sid,application['id'],role['current_version'],uid(),'COMPLETED',dump(result),now()))
    service.db.execute("UPDATE applications SET assessment_id=?,status='ASSESSED' WHERE id=?",(sid,application['id']))
    return source


def test_excel_separates_model_interpretation_and_keeps_entire_long_source(service,application,role):
    source=seed_score(service,application,role)
    before=service.db.one('SELECT * FROM followups WHERE application_id=?',(application['id'],))
    output=overview_export(service.db,service.storage)
    book=load_workbook(service.storage.file(output['id'])[1])
    sheet=book['评估证据']
    assert sheet['I1'].value=='简历原文' and sheet['K1'].value=='模型解读'
    assert ''.join(sheet.cell(i,9).value or '' for i in range(2,sheet.max_row+1))==source
    assert sheet['I2'].data_type=='s' and sheet['K2'].value=='模型对相关能力的概括 <解读>'
    assert len(book['招聘推荐总览'][1])==18
    assert service.db.one('SELECT * FROM followups WHERE application_id=?',(application['id'],))==before


def test_html_separates_semantic_summary_from_escaped_original_source(service,application,role):
    seed_score(service,application,role)
    output=role_reports(service.db,service.storage,uid())[0]
    text=service.storage.file(output['id'])[1].read_text(encoding='utf-8')
    assert '模型解读：模型对相关能力的概括 &lt;解读&gt;' in text
    assert '<summary>对应原文 R1</summary>' in text
    assert '原始资料 &lt;原文&gt;' in text
    assert '<解读>' not in text and '<原文>' not in text


def test_legacy_reports_keep_existing_quotes_without_claiming_model_review(service,application,role):
    seed_score(service,application,role,False)
    output=role_reports(service.db,service.storage,uid())[0]
    text=service.storage.file(output['id'])[1].read_text(encoding='utf-8')
    assert '原始资料 &lt;原文&gt;' in text
    assert '模型解读：' not in text and '分析与语义复核均由模型完成' not in text
