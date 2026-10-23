from datetime import datetime
from io import BytesIO
import json
import pytest
from openpyxl import load_workbook

from hr_assist.reports import followup_export, followup_import, overview_export, applications
from hr_assist.storage import BusinessError


def exported(service):
    record=followup_export(service.db,service.storage)
    return service.storage.file(record['id'])[1].read_bytes()


def edit(data,changes):
    book=load_workbook(BytesIO(data));sheet=book['人工跟进']
    for position,value in changes.items():sheet[position]=value
    stream=BytesIO();book.save(stream);return stream.getvalue()


def upload(service,data):return followup_import(service.db,service.storage,data,'人工跟进.xlsx')


def test_real_excel_roundtrip_formulas_dates_and_overview(service,application):
    data=exported(service)
    assert upload(service,data)['updated_rows']==0
    edited=edit(data,{'D2':datetime(2026,10,23),'F2':'通过','G2':'=1+1','H2':'安排面试'})
    assert upload(service,edited)['updated_rows']==1
    assert upload(service,edited)['updated_rows']==0
    record=overview_export(service.db,service.storage)
    book=load_workbook(service.storage.file(record['id'])[1],data_only=False)
    sheet=book['招聘推荐总览']
    assert sheet['K2'].value==datetime(2026,10,23)
    assert sheet['M2'].value=='通过' and sheet['N2'].value=='=1+1' and sheet['N2'].data_type=='f'
    assert sheet['O2'].value=='安排面试'
    assert sheet.column_dimensions['P'].hidden
    assert len(sheet[1])==16
    assert json.loads(service.db.one('SELECT cells FROM followups')['cells'])[3]['value']=='=1+1'


def test_stale_nonoverlapping_edits_merge_conflicts_rollback(service,application):
    baseline=exported(service)
    upload(service,edit(baseline,{'H2':'待面试'}))
    upload(service,edit(baseline,{'F2':'通过'}))
    assert applications(service.db)[0]['manual']['当前状态']=='待面试'
    with pytest.raises(BusinessError,match='冲突'):
        upload(service,edit(baseline,{'H2':'放弃','G2':'不可写入'}))
    row=applications(service.db)[0]
    assert row['manual']['当前状态']=='待面试'
    assert row['manual']['面试评价'] is None
    assert upload(service,edit(exported(service),{'H2':None}))['updated_rows']==1
    assert applications(service.db)[0]['manual']['当前状态'] is None


def test_omitted_rows_preserved_invalid_ids_rejected(service,application):
    baseline=exported(service)
    upload(service,edit(baseline,{'H2':'已联系'}))
    book=load_workbook(BytesIO(baseline));book['人工跟进'].delete_rows(2)
    stream=BytesIO();book.save(stream)
    assert upload(service,stream.getvalue())['omitted_rows_preserved']==1
    assert applications(service.db)[0]['manual']['当前状态']=='已联系'
    with pytest.raises(BusinessError,match='记录 ID'):
        upload(service,edit(baseline,{'A2':'malicious-id'}))


def test_machine_text_formula_injection(service,application):
    service.db.execute("UPDATE applications SET name='=HYPERLINK(\"https://example.invalid\")' WHERE id=?",(application['id'],))
    file=overview_export(service.db,service.storage)
    sheet=load_workbook(service.storage.file(file['id'])[1])['招聘推荐总览']
    assert sheet['C2'].data_type=='s'
