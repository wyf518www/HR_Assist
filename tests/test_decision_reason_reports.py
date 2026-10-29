import pytest
from openpyxl import load_workbook

from hr_assist.db import dump, now, uid
from hr_assist.reports import OVERVIEW_COLUMNS, empty_cells, overview_export, role_reports


def seed(service, application, role, result):
    identity = uid()
    service.db.execute('INSERT INTO assessments VALUES (?,?,?,?,?,?,?)',
                       (identity, application['id'], role['current_version'], uid(), 'COMPLETED', dump(result), now()))
    service.db.execute("UPDATE applications SET assessment_id=?,status='ASSESSED' WHERE id=?", (identity, application['id']))
    cells = empty_cells()
    cells[3] = {'value': '=1+1', 'type': 'f', 'format': 'General'}
    cells[4] = {'value': '原人工备注', 'type': 's', 'format': '@'}
    cells[5] = {'value': '已面', 'type': 's', 'format': 'General'}
    service.db.execute('UPDATE followups SET cells=?,revision=3 WHERE application_id=?', (dump(cells), application['id']))


def result_with_reasons():
    return {'score': 92, 'grade': '待定', 'base_grade': '强烈推荐', 'grade_limited': True,
            'provisional': True, 'gate_reason': '保存的学历政策说明', 'facts': {}, 'items': [],
            'advantages': ['原优势'], 'risks': ['原风险'], 'policy_version': 'admission-default-v1',
            'decision_reasons': [
                {'code': 'SYNTHETIC', 'title': '=合成原因 <标题>', 'detail': '需要核对 <具体说明>', 'action': '核对来源 <处理>'},
                {'code': 'CRITERION_UNKNOWN', 'title': '年限尚待核验', 'detail': '任职年月缺失', 'action': '补齐任职年月', 'criterion_id': 'experience'}]}


def test_overview_adds_all_reasons_and_actions_without_changing_columns_or_human_content(service, application, role):
    result = result_with_reasons()
    seed(service, application, role, result)
    stored = service.db.all('SELECT * FROM assessments')
    followup = service.db.one('SELECT * FROM followups WHERE application_id=?', (application['id'],))
    output = overview_export(service.db, service.storage)
    book = load_workbook(service.storage.file(output['id'])[1], data_only=False)
    sheet = book['招聘推荐总览']
    assert [cell.value for cell in sheet[1]] == [*OVERVIEW_COLUMNS, '_application_id']
    assert sheet['D2'].value == 92 and sheet['E2'].value == '待定'
    assert sheet['I2'].data_type == sheet['J2'].data_type == 's'
    for reason in result['decision_reasons']:
        assert reason['title'] in sheet['I2'].value and reason['detail'] in sheet['I2'].value
        assert reason['action'] in sheet['J2'].value
    assert '原风险' in sheet['I2'].value and sheet['H2'].value == '原优势'
    assert '按分数建议等级：强烈推荐' in sheet['J2'].value
    assert '最终等级调整为待定' in sheet['J2'].value
    assert sheet['N2'].value == '=1+1' and sheet['N2'].data_type == 'f'
    assert sheet['O2'].value == '原人工备注' and sheet['P2'].value == '已面'
    assert sheet.column_dimensions['R'].hidden
    assert service.db.all('SELECT * FROM assessments') == stored
    assert service.db.one('SELECT * FROM followups WHERE application_id=?', (application['id'],)) == followup


def test_html_shows_grade_limit_and_escaped_reason_detail_and_action(service, application, role):
    result = result_with_reasons()
    seed(service, application, role, result)
    output = role_reports(service.db, service.storage, uid())[0]
    html = service.storage.file(output['id'])[1].read_text(encoding='utf-8')
    assert '<h3>为何待定</h3>' in html
    assert '按分数建议等级：<b>强烈推荐</b>' in html
    assert '最终等级调整为<b>待定</b>' in html
    assert '=合成原因 &lt;标题&gt;' in html
    assert '需要核对 &lt;具体说明&gt;' in html
    assert '核对来源 &lt;处理&gt;' in html
    assert '年限尚待核验' in html and '补齐任职年月' in html
    assert all(raw not in html for raw in ('<标题>', '<具体说明>', '<处理>'))


@pytest.mark.parametrize('with_saved_education', [False, True])
def test_history_explanations_are_read_only_and_do_not_apply_the_new_policy(service, application, role, with_saved_education):
    result = {'score': 92, 'grade': '待定', 'provisional': True, 'gate_reason': '历史学历条件不明', 'facts': {}, 'items': []}
    if with_saved_education:
        result.update(education_gate='unknown', education={
            'degree': 'bachelor', 'completion': 'attained', 'admission': 'unknown', 'institution_match': 'catalog',
            'evidence': [{'span_id': 'R1', 'quote': '历史本科资料', 'source': 'text'}]})
    seed(service, application, role, result)
    before = service.db.all('SELECT * FROM assessments')
    report = overview_export(service.db, service.storage)
    sheet = load_workbook(service.storage.file(report['id'])[1])['招聘推荐总览']
    assert sheet['D2'].value == 92 and sheet['E2'].value == '待定'
    assert '重新评估' in sheet['J2'].value
    assert '按分数建议等级' not in sheet['J2'].value
    if with_saved_education:
        assert '历史评估未确认本科招生方式' in sheet['I2'].value
    else:
        assert '无法准确还原具体缺项' in sheet['I2'].value
    html_file = role_reports(service.db, service.storage, uid())[0]
    html = service.storage.file(html_file['id'])[1].read_text(encoding='utf-8')
    assert '<h3>为何待定</h3>' in html and '重新评估' in html
    assert '按分数建议等级' not in html
    assert service.db.all('SELECT * FROM assessments') == before


def test_score_band_reason_is_reported_without_claiming_a_limit(service, application, role):
    result = {'score': 65, 'grade': '待定', 'base_grade': '待定', 'grade_limited': False,
              'gate_reason': '', 'provisional': False, 'facts': {}, 'items': []}
    seed(service, application, role, result)
    output = role_reports(service.db, service.storage, uid())[0]
    html = service.storage.file(output['id'])[1].read_text(encoding='utf-8')
    assert '分数位于待定区间' in html and '65' in html
    assert '最终等级调整' not in html
